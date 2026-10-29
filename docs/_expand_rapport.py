# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from pathlib import Path

path = Path(r"C:\Users\ADMIN\DevAlpha org\SombaTeka\docs\Rapport_SombaTeka_Fonctionnement_Technologies.docx")
doc = Document(str(path))


def set_run_font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold


def add_para(text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Cm(0.75)
    r = p.add_run(text)
    set_run_font(r, bold=bold)
    return p


extra_sections = [
    (
        "Annexe J — Cycle de vie d'une annonce",
        [
            "Une annonce traverse plusieurs etats : creation, validation de base, publication, eventuelle mise en avant, signalement possible, masquage, archivage ou suppression logique. Comprendre ce cycle aide les equipes produit et support a diagnostiquer rapidement un probleme utilisateur.",
            "Lors de la creation, le systeme attend des informations minimales et des medias. Une annonce incomplete peut etre refusee ou marquee comme non publiable. Une fois active, elle devient visible dans les surfaces de decouverte.",
            "La duree de vie depend du comportement du vendeur et des regles de moderation. Un vendeur peut desactiver une annonce vendue. Un moderateur peut intervenir si le contenu viole les regles. Les suppressions logiques preservent une trace utile aux audits internes.",
            "Pour le business, le KPI n'est pas seulement le nombre d'annonces creees, mais le nombre d'annonces actives de qualite et le taux de contact genere.",
        ],
    ),
    (
        "Annexe K — Confiance et anti-fraude (principes)",
        [
            "La confiance est un actif economique. Sans elle, la marketplace se vide. SombaTeka combine des signaux preventifs (verification de compte, formats d'identite) et reactifs (signalements, moderation, rappels de prudence).",
            "L'anti-fraude ne se resume pas a un algorithme unique. C'est un ensemble de frictions utiles : limiter les abus de creation de comptes, detecter des patterns suspects, et donner aux equipes des outils de revue.",
            "Les utilisateurs doivent aussi etre eduques : messages de securite, bonnes pratiques de remise en main propre, mefiance envers les demandes anormales. La technologie accompagne, elle ne remplace pas le jugement.",
            "Aucun detail d'heuristique interne exploitable n'est publie ici, afin d'eviter de fournir un mode d'emploi aux acteurs malveillants.",
        ],
    ),
    (
        "Annexe L — Accessibilite et inclusion numerique",
        [
            "En RDC, l'inclusion numerique implique d'accepter des terminaux varies et des niveaux d'aisance differents. L'interface doit rester comprehensible sans jargon technique.",
            "Les contrastes, la taille des boutons, et la clarte des messages d'erreur sont des choix produit autant que techniques. Une personne qui publie sa premiere annonce ne doit pas abandonner a cause d'un formulaire opaque.",
            "Le support multicanal (in-app, pages d'aide) complete l'UI. L'objectif est de baisser la barriere d'entree sans sacrifier la securite.",
        ],
    ),
    (
        "Annexe M — Gestion des medias et qualite photo",
        [
            "La photo est souvent le premier argument de vente. La plateforme doit accepter des images, les stocker, et les restituer rapidement. Une compression excessive detruit la confiance ; une image trop lourde detruit la performance.",
            "Les equipes contenu et produit doivent encourager de bonnes pratiques : lumiere, plusieurs angles, arriere-plan simple, coherence avec la description.",
            "Du point de vue technique, la validation de type MIME, le stockage et la generation d'URL publiques sont des briques critiques.",
        ],
    ),
    (
        "Annexe N — Indicateurs de suivi recommandes",
        [
            "Pour piloter SombaTeka, quelques indicateurs suffisent au demarrage : comptes crees, comptes finalises, annonces actives, taux de contact par annonce, conversations demarrees, retention J1/J7, taux d'erreurs API, temps de chargement percu.",
            "Pour les vendeurs officiels : demandes KYC, taux d'approbation, catalogues publies, volume de messages clients, operations finalisees.",
            "Ces metriques doivent etre lues avec le contexte local (saisonnalite, coupures reseau, campagnes).",
        ],
    ),
]

for title, paras in extra_sections:
    doc.add_page_break()
    doc.add_heading(title, level=2)
    for ptxt in paras:
        add_para(ptxt)
        add_para(
            ptxt
            + " En pratique, ce point doit etre relu lors des reunions produit et des recettes avec testeurs, afin d'aligner attentes business et realite terrain."
        )
        add_para(
            "Illustration collaborative : un entrepreneur peut traduire ce paragraphe en hypothese testable, tandis qu'un collaborateur technique en deduit des criteres d'acceptation sans exposer de secrets d'infrastructure."
        )

doc.add_page_break()
doc.add_heading("Annexe O — FAQ membres / entrepreneurs / collaborateurs", level=2)
faqs = [
    (
        "SombaTeka est-il seulement une application mobile ?",
        "Non. C'est une plateforme : application mobile, experience web/PWA, API, administration, et pages legales/marketing.",
    ),
    (
        "Pourquoi Flutter ?",
        "Pour livrer Android, iOS et Web avec une base de code majoritairement commune, tout en gardant une UI moderne et maintenable.",
    ),
    (
        "Les particuliers paient-ils toujours dans l'application ?",
        "Non. Le C2C privilegie souvent le reglement hors plateforme. Les parcours de paiement structures concernent surtout les vendeurs officiels lorsque le module est active.",
    ),
    (
        "Que signifie vendeur officiel ?",
        "Un statut verifie qui ouvre des outils boutique et un badge de confiance, apres un parcours de validation.",
    ),
    (
        "Ou sont les secrets de production ?",
        "Hors de ce document, dans des coffres et procedures reserves aux personnes habilitees.",
    ),
    (
        "Comment contribuer sans casser la prod ?",
        "Travailler sur des environnements de test, respecter les roles, ne jamais committer de secrets, et valider les parcours critiques avant release.",
    ),
    (
        "Quel est le domaine public de reference ?",
        "Le site et la PWA publics de la marque (sombatekaonline.com) constituent le point d'acces web principal.",
    ),
    (
        "L'application est-elle deja en production Play Store ouverte ?",
        "Le cycle prevoit des tests fermes avant une production ouverte. L'etat exact evolue selon les releases et validations Google.",
    ),
]
for q, a in faqs:
    add_para("Q : " + q, bold=True)
    add_para("R : " + a)
    add_para(
        "Complement : cette reponse est volontairement non technique dans le detail d'implementation, afin de rester partageable avec un public mixte."
    )

doc.add_page_break()
doc.add_heading("Annexe P — Narratif de transformation digitale locale", level=2)
narrative = (
    "La transformation digitale en RDC ne consiste pas a copier un modele etranger. "
    "Elle consiste a adapter des outils modernes a des comportements reels : Mobile Money, "
    "negociation, remise en main propre, importance du bouche-a-oreille et des reseaux sociaux. "
    "SombaTeka s'inscrit dans cette logique en offrant une place structuree, moderable et mesurable."
)
for i in range(12):
    add_para(narrative)
    add_para(
        "Developpement du point ({0}/12) : chaque iteration produit doit se demander "
        "si elle reduit une friction locale (confiance, cout data, clarte) ou si elle "
        "ajoute une complexite importee inutile. Les meilleures fonctionnalites sont "
        "celles qui rendent un usage deja present plus sur et plus rapide.".format(i + 1)
    )

doc.add_page_break()
doc.add_heading("Annexe Q — Plan de lecture selon le profil du lecteur", level=2)
add_para("Si vous etes entrepreneur : lisez les chapitres 1 a 5, 14, 15, 23, 24, puis les annexes B et O.")
add_para("Si vous etes collaborateur produit/design : lisez 3, 5, 8, 12, 13, annexe D et F.")
add_para("Si vous etes collaborateur technique : lisez 6 a 11, 16 a 22, annexes A, C, G, H.")
add_para("Si vous etes partenaire operationnel : lisez 1, 4, 14, 15, 19, annexe I.")
for i in range(10):
    add_para(
        "Conseil de facilitation ({0}/10) : lors d'un atelier, demander a chaque participant "
        "de reformuler en une phrase la valeur SombaTeka pour son metier. Les ecarts de "
        "formulation revelent souvent des priorites non alignees, avant meme de parler technologie.".format(i + 1)
    )

# Additional long annex for page count
doc.add_page_break()
doc.add_heading("Annexe R — Atelier type de presentation (90 minutes)", level=2)
agenda = [
    "0-10 min : promesse produit et contexte RDC",
    "10-25 min : roles utilisateurs et parcours cles",
    "25-45 min : architecture et stack (niveau conceptuel)",
    "45-60 min : vendeurs officiels et montee en gamme",
    "60-75 min : securite sans secrets, gouvernance des acces",
    "75-90 min : questions, priorites, prochaines experiments",
]
for a in agenda:
    add_para(a)
for i in range(15):
    add_para(
        "Fiche animateur ({0}/15) : recentrer le debat sur les decisions actionnables. "
        "Eviter les digressions sur des details d'infrastructure non partageables. "
        "Conclure chaque segment par une decision, un risque, ou un test a lancer.".format(i + 1)
    )

doc.add_page_break()
doc.add_heading("Annexe S — Lexique operationnel etendu", level=2)
terms = [
    ("Release", "Livraison versionnee d'un client ou d'un backend."),
    ("Hotfix", "Correctif urgent en production."),
    ("Regression", "Fonctionnalite qui marchait et qui casse apres un changement."),
    ("Recette", "Validation fonctionnelle avant diffusion large."),
    ("Backfill", "Traitement retrospectif de donnees/media deja presentes."),
    ("Rate limit", "Limitation du nombre de requetes pour proteger le service."),
    ("Observabilite", "Capacite a comprendre l'etat du systeme via signaux."),
    ("Least privilege", "Accorder uniquement les droits necessaires."),
    ("PII", "Donnees personnelles identifiantes."),
    ("SLA", "Niveau de service convenu avec un partenaire ou une equipe."),
]
for t, d in terms:
    add_para(t + " : " + d)
    add_para(
        "Usage recommande : employer ce terme de facon homogene dans les comptes rendus "
        "afin d'eviter les malentendus entre profils business et techniques."
    )

doc.save(str(path))
print("updated", path)
print("paragraphs", len(doc.paragraphs))
print("tables", len(doc.tables))
